"""Export the claim-verification SQLite table definition to an Excel workbook."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Data_List_Claim_Verification" / "claim_verification_table_definition.xlsx"

TABLES = {
    "dataset_registry": ("적재 원본 관리", [
        ("dataset_key", "TEXT", "PK", "데이터셋 식별자"),
        ("source_file", "TEXT", "", "원본 파일명"),
        ("description", "TEXT", "", "데이터 설명"),
        ("record_count", "INTEGER", "", "적재 레코드 수"),
        ("loaded_at", "TEXT", "", "적재 시각(UTC)"),
    ]),
    "products": ("건강기능식품 제품정보", [
        ("product_id", "INTEGER", "PK", "내부 제품 식별자"),
        ("manufacturing_report_number", "TEXT", "Index", "제조신고번호"),
        ("product_name", "TEXT", "", "제품명 원문"),
        ("product_name_normalized", "TEXT", "Index", "검색용 정규화 제품명"),
        ("company_name", "TEXT", "", "제조사·업체명"),
        ("registration_date", "TEXT", "", "등록일"),
        ("main_functionality", "TEXT", "", "원본 기능성 내용"),
        ("intake_method", "TEXT", "", "섭취방법"),
        ("intake_precautions", "TEXT", "", "섭취 시 주의사항"),
        ("specification", "TEXT", "", "제품 규격"),
        ("appearance", "TEXT", "", "성상"),
        ("shelf_life", "TEXT", "", "유통기한"),
        ("storage_and_distribution_standard", "TEXT", "", "보관·유통 기준"),
    ]),
    "ingredients": ("기능성·개별인정형 원료", [
        ("ingredient_id", "INTEGER", "PK", "내부 원료 식별자"),
        ("ingredient_type", "TEXT", "", "FUNCTIONAL_INGREDIENT 또는 INDIVIDUALLY_APPROVED"),
        ("recognition_number", "TEXT", "Index", "원료 인정번호"),
        ("ingredient_name", "TEXT", "", "원료명 원문"),
        ("ingredient_name_normalized", "TEXT", "Index", "검색용 정규화 원료명"),
        ("functionality_text", "TEXT", "", "공식 기능성·주기능성 내용"),
        ("daily_intake_text", "TEXT", "", "일일 섭취량"),
        ("daily_intake_min", "TEXT", "", "최소 섭취량"),
        ("daily_intake_max", "TEXT", "", "최대 섭취량"),
        ("intake_precautions", "TEXT", "", "섭취 시 주의사항"),
        ("company_name", "TEXT", "", "신청·업체명"),
        ("approval_date", "TEXT", "", "인정·허가일"),
        ("company_address", "TEXT", "", "업체 주소"),
    ]),
    "diseases": ("질병 코드 사전", [
        ("disease_id", "INTEGER", "PK", "내부 질병 식별자"),
        ("disease_code", "TEXT", "Index", "표준 질병 코드"),
        ("disease_name", "TEXT", "", "국문 질병명"),
        ("disease_name_normalized", "TEXT", "Index", "검색용 정규화 질병명"),
        ("disease_name_en", "TEXT", "", "영문 질병명"),
        ("sick_type", "TEXT", "Unique 조합", "상병 구분 코드"),
        ("sick_type_label", "TEXT", "", "상병 구분명"),
        ("medicine_type", "TEXT", "Unique 조합", "의료 구분 코드"),
        ("medicine_type_label", "TEXT", "", "의료 구분명"),
    ]),
    "claims": ("검증 대상 claim", [
        ("claim_id", "INTEGER", "PK", "claim 식별자"),
        ("claim_origin", "TEXT", "", "VIDEO, MANUAL"),
        ("product_id", "INTEGER", "FK → products", "관련 제품"),
        ("claim_text", "TEXT", "", "검증할 원문 문장"),
        ("claim_type", "TEXT", "", "claim 분류(예: FUNCTIONALITY)"),
        ("subject_text", "TEXT", "", "claim 주체(추후 추출)"),
        ("predicate_text", "TEXT", "", "claim 관계·효능(추후 추출)"),
        ("target_text", "TEXT", "", "claim 대상(추후 추출)"),
        ("disease_id", "INTEGER", "FK → diseases", "연결된 질병(추후 매핑)"),
        ("extracted_at", "TEXT", "", "claim 생성·추출 시각"),
    ]),
    "evidence": ("공식 검증 근거", [
        ("evidence_id", "INTEGER", "PK", "근거 식별자"),
        ("ingredient_id", "INTEGER", "FK → ingredients", "근거 대상 원료"),
        ("evidence_type", "TEXT", "", "OFFICIAL_FUNCTIONALITY, OFFICIAL_INTAKE, OFFICIAL_PRECAUTION"),
        ("evidence_text", "TEXT", "", "공식 근거 문장"),
        ("source_authority", "TEXT", "", "현재 MFDS"),
    ]),
    "claim_evidence_links": ("claim 사실성 판정", [
        ("link_id", "INTEGER", "PK", "연결 식별자"),
        ("claim_id", "INTEGER", "FK → claims", "검증 대상 claim"),
        ("evidence_id", "INTEGER", "FK → evidence", "비교한 공식 근거"),
        ("match_method", "TEXT", "", "EXACT, RULE, LLM, MANUAL"),
        ("verdict", "TEXT", "", "SUPPORTED, CONTRADICTED, INSUFFICIENT, REVIEW_REQUIRED"),
        ("confidence", "REAL", "", "0~1 신뢰도"),
        ("reason", "TEXT", "", "판정 사유"),
        ("created_at", "TEXT", "", "연결·판정 시각"),
    ]),
}


def style_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for index, width in enumerate((30, 15, 20, 58), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main():
    workbook = Workbook()
    summary = workbook.active
    summary.title = "테이블 요약"
    summary.append(["테이블", "역할", "컬럼 수"])
    for table_name, (description, columns) in TABLES.items():
        summary.append([table_name, description, len(columns)])
    style_sheet(summary)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 42

    for table_name, (description, columns) in TABLES.items():
        sheet = workbook.create_sheet(table_name[:31])
        sheet.append(["컬럼명", "타입", "키·제약", "설명"])
        for column in columns:
            sheet.append(column)
        style_sheet(sheet)
        sheet.sheet_properties.tabColor = "1F4E78"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
